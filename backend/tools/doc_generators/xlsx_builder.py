import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class XLSXDeliverableBuilder:
    """Generates professional, styled Excel (.xlsx) workbooks for engineering calculations and procurement spreadsheets."""

    @staticmethod
    def create_spreadsheet(title: str, headers: list[str], rows: list[list], summary_data: dict, output_path: str) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Industrial Synthesis"

        # Gridlines enabled
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="0F2043")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="555555")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F2043", end_color="0F2043", fill_type="solid")
        
        row_even_fill = PatternFill(start_color="F4F7FA", end_color="F4F7FA", fill_type="solid")
        row_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        summary_fill = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
        summary_font = Font(name="Calibri", size=11, bold=True, color="0F2043")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # Title Block
        ws.merge_cells("A1:G1")
        cell_title = ws["A1"]
        cell_title.value = title.upper()
        cell_title.font = title_font
        cell_title.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:G2")
        cell_sub = ws["A2"]
        cell_sub.value = "MRPL Sovereign Workbench Engine | Air-Gapped Engineering Synthesis Sheet"
        cell_sub.font = subtitle_font
        ws.row_dimensions[2].height = 18

        # Blank row 3
        ws.row_dimensions[3].height = 10

        # Table Header Row (Row 4)
        start_row = 4
        ws.row_dimensions[start_row].height = 25
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Data Rows (Row 5+)
        current_row = start_row + 1
        for row_idx, row_values in enumerate(rows):
            ws.row_dimensions[current_row].height = 20
            fill = row_even_fill if row_idx % 2 == 1 else row_odd_fill
            for col_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.value = val
                cell.fill = fill
                cell.border = thin_border
                
                # Alignments
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, float):
                        cell.number_format = '#,##0.00'
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        # Summary Metrics Block below table
        if summary_data:
            current_row += 1
            ws.cell(row=current_row, column=1).value = "SUMMARY ANALYSIS & METRICS"
            ws.cell(row=current_row, column=1).font = summary_font
            ws.row_dimensions[current_row].height = 22
            current_row += 1

            for key, val in summary_data.items():
                ws.row_dimensions[current_row].height = 20
                c_k = ws.cell(row=current_row, column=1)
                c_k.value = key
                c_k.font = Font(name="Calibri", size=10, bold=True)
                c_k.fill = summary_fill
                c_k.border = thin_border

                c_v = ws.cell(row=current_row, column=2)
                c_v.value = val
                c_v.font = Font(name="Calibri", size=10, bold=True, color="0F2043")
                c_v.fill = summary_fill
                c_v.border = thin_border
                if isinstance(val, (int, float)):
                    c_v.alignment = Alignment(horizontal="right", vertical="center")
                current_row += 1

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value and cell.row > 2:  # skip merged title rows
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # Save workbook
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return str(path)

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

class PPTXDeliverableBuilder:
    """Upgraded enterprise PowerPoint (.pptx) deck generator for ODIN Sovereign Workbench."""

    @staticmethod
    def create_presentation(title: str, subtitle: str, slides_data: list[dict], output_path: str) -> str:
        prs = Presentation()
        
        # 16:9 Widescreen Layout
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]
        slide_title = prs.slides.add_slide(blank_layout)

        # Dark Metallic Background Box
        bg_shape = slide_title.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
        bg_shape.line.fill.background()

        # Title Text Box
        tb_title = slide_title.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.5))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True

        p1 = tf_title.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(255, 255, 255)
        p1.font.name = "Calibri"

        p2 = tf_title.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(148, 163, 184)
        p2.font.name = "Calibri"
        p2.space_before = Pt(12)

        p3 = tf_title.add_paragraph()
        p3.text = "ODIN — On-Premise Data & Industrial Intelligence Workbench | MRPL SIH26117"
        p3.font.size = Pt(12)
        p3.font.bold = True
        p3.font.color.rgb = RGBColor(56, 189, 248) # Sky Blue Accent
        p3.font.name = "Calibri"
        p3.space_before = Pt(28)

        # Dynamic Content Slides
        for slide_info in slides_data:
            s_heading = slide_info.get("heading", "Executive Summary")
            s_bullets = slide_info.get("bullets", [])
            s_metrics = slide_info.get("metrics", None)

            slide = prs.slides.add_slide(blank_layout)

            # Top Header Bar
            header_box = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
            header_box.fill.solid()
            header_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
            header_box.line.fill.background()

            tb_head = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.8))
            p_head = tb_head.text_frame.paragraphs[0]
            p_head.text = s_heading
            p_head.font.size = Pt(24)
            p_head.font.bold = True
            p_head.font.color.rgb = RGBColor(255, 255, 255)
            p_head.font.name = "Calibri"

            # Bullets Container
            tb_content = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.2))
            tf_content = tb_content.text_frame
            tf_content.word_wrap = True

            for idx, bullet in enumerate(s_bullets):
                p_b = tf_content.add_paragraph() if idx > 0 else tf_content.paragraphs[0]
                p_b.text = f"•  {bullet}"
                p_b.font.size = Pt(16)
                p_b.font.color.rgb = RGBColor(30, 41, 59)
                p_b.font.name = "Calibri"
                p_b.space_after = Pt(14)

            # Metrics Cards Column
            if s_metrics and isinstance(s_metrics, list):
                card_y = Inches(1.5)
                for metric in s_metrics[:3]:
                    m_label = metric.get("label", "Metric")
                    m_val = metric.get("value", "N/A")

                    card = slide.shapes.add_shape(1, Inches(8.7), card_y, Inches(3.8), Inches(1.4))
                    card.fill.solid()
                    card.fill.fore_color.rgb = RGBColor(241, 245, 249)
                    card.line.color.rgb = RGBColor(203, 213, 225)

                    tb_m = slide.shapes.add_textbox(Inches(8.8), card_y + Inches(0.15), Inches(3.6), Inches(1.1))
                    p_mv = tb_m.text_frame.paragraphs[0]
                    p_mv.text = str(m_val)
                    p_mv.font.size = Pt(24)
                    p_mv.font.bold = True
                    p_mv.font.color.rgb = RGBColor(15, 23, 42)
                    p_mv.alignment = PP_ALIGN.CENTER

                    p_ml = tb_m.text_frame.add_paragraph()
                    p_ml.text = str(m_label).upper()
                    p_ml.font.size = Pt(10)
                    p_ml.font.bold = True
                    p_ml.font.color.rgb = RGBColor(100, 116, 139)
                    p_ml.alignment = PP_ALIGN.CENTER

                    card_y += Inches(1.6)

        # Save Presentation
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(path))
        return str(path)
